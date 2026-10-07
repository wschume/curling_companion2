#!/usr/bin/env python3
"""Validate, preview, and atomically import an Excel tournament list."""
import argparse
from datetime import date, datetime
import hashlib
import json
import math
from pathlib import Path
import re
import sys
import unicodedata

from openpyxl import Workbook, load_workbook

REQUIRED = ('name', 'city', 'country', 'club', 'contactInformation', 'startDate', 'endDate')
OPTIONAL = ('signupDeadline', 'entryFee', 'currency', 'maxNumberOfTeams', 'websiteUrl')
FIELDS = REQUIRED + OPTIONAL
CURRENCIES = {'EUR', 'USD', 'GBP', 'CHF', 'NOK'}


class ValidationError(ValueError):
    pass


def normalize(value):
    return ' '.join(unicodedata.normalize('NFKC', value).casefold().split())


def calendar_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
        return date.fromisoformat(value)
    raise ValueError('expected Excel date or YYYY-MM-DD')


def identity(record):
    # Existing app documents use local ISO strings, sometimes with time components.
    start = record['startDate']
    if isinstance(start, str):
        start = start[:10]
    return (normalize(record['name']), normalize(record['city']), calendar_date(start).isoformat())


def document(record, organizer_id):
    data = {field: record.get(field) for field in FIELDS}
    for field in ('startDate', 'endDate', 'signupDeadline'):
        if data[field] is not None:
            data[field] = calendar_date(data[field]).isoformat() + 'T00:00:00.000Z'
    data['organizerId'] = organizer_id
    data['id'] = hashlib.sha256(json.dumps(identity(record), ensure_ascii=False, separators=(',', ':')).encode()).hexdigest()
    return data


def template(path):
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Tournaments'
    sheet.append(FIELDS)
    sheet.freeze_panes = 'A2'
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 24
    instructions = wb.create_sheet('Instructions')
    for line in (
        'Required columns: ' + ', '.join(REQUIRED),
        'Dates: Excel dates or YYYY-MM-DD. End must follow start by at most five days.',
        'Signup deadline, if present, must precede start. Contact: email or phone (7+ digits).',
        'Fees and team counts must be nonnegative; team counts must be integers.',
        'Currencies: EUR, USD, GBP, CHF, NOK. Blank currency defaults to EUR with a fee.',
        'Maximum 500 tournaments. No formulas. Duplicates block the entire import.',
    ):
        instructions.append([line])
    wb.save(path)
    wb.close()


def parse(path):
    errors, records = [], []
    wb = load_workbook(path, data_only=False)
    try:
        if 'Tournaments' not in wb.sheetnames:
            raise ValidationError('Missing Tournaments worksheet')
        rows = iter(wb['Tournaments'].iter_rows())
        header_cells = next(rows, ())
        headers = [str(c.value).strip() if c.value is not None else '' for c in header_cells]
        # Ignore unused trailing columns, but not columns containing data below them.
        while headers and not headers[-1]:
            headers.pop()
        if len(set(headers)) != len(headers):
            errors.append('Row 1: duplicate headers')
        for field in REQUIRED:
            if field not in headers:
                errors.append(f'Row 1: missing header {field}')
        for header in headers:
            if header not in FIELDS:
                errors.append(f'Row 1: unknown header {header!r}')
        if any(c.data_type == 'f' for c in header_cells):
            errors.append('Row 1: formulas are not allowed')
        if errors:
            raise ValidationError('\n'.join(errors))
        seen = {}
        for row_number, cells in enumerate(rows, 2):
            if all(c.value is None or c.value == '' for c in cells):
                continue
            before = len(errors)
            record = {'_row': row_number}
            for index, cell in enumerate(cells):
                field = headers[index] if index < len(headers) else f'column {index + 1}'
                if cell.data_type == 'f':
                    errors.append(f'Row {row_number}, {field}: formulas are not allowed')
                if index >= len(headers):
                    if cell.value is not None:
                        errors.append(f'Row {row_number}, {field}: data without header')
                    continue
                value = cell.value.strip() if isinstance(cell.value, str) else cell.value
                record[field] = None if value == '' else value
            for field in REQUIRED:
                if record.get(field) is None:
                    errors.append(f'Row {row_number}, {field}: required')
            for field in ('name', 'city', 'country', 'club', 'contactInformation', 'websiteUrl', 'currency'):
                if record.get(field) is not None and not isinstance(record[field], str):
                    errors.append(f'Row {row_number}, {field}: expected text')
            for field in ('startDate', 'endDate', 'signupDeadline'):
                if record.get(field) is not None:
                    try:
                        record[field] = calendar_date(record[field])
                    except ValueError as exc:
                        errors.append(f'Row {row_number}, {field}: {exc}')
                        record[field] = None
            start, end, deadline = (record.get(f) for f in ('startDate', 'endDate', 'signupDeadline'))
            if start and end and not 0 < (end - start).days <= 5:
                errors.append(f'Row {row_number}, endDate: must follow start by at most five days')
            if deadline and start and deadline >= start:
                errors.append(f'Row {row_number}, signupDeadline: must precede start')
            contact = record.get('contactInformation')
            if isinstance(contact, str) and not (re.fullmatch(r'[^@\s]+@[^@\s]+\.[^@\s]+', contact) or len(re.sub(r'[^0-9]', '', contact)) >= 7):
                errors.append(f'Row {row_number}, contactInformation: expected email or phone')
            for field in ('entryFee', 'maxNumberOfTeams'):
                value = record.get(field)
                if value is not None:
                    try:
                        if isinstance(value, bool):
                            raise ValueError()
                        number = float(value)
                        if not math.isfinite(number) or number < 0 or (field == 'maxNumberOfTeams' and not number.is_integer()):
                            raise ValueError()
                        record[field] = int(number) if field == 'maxNumberOfTeams' else number
                    except (ValueError, TypeError, OverflowError):
                        errors.append(f'Row {row_number}, {field}: expected finite nonnegative ' + ('integer' if field == 'maxNumberOfTeams' else 'number'))
            currency = record.get('currency')
            if currency is not None and currency not in CURRENCIES:
                errors.append(f'Row {row_number}, currency: unsupported currency')
            if record.get('entryFee') is not None and currency is None:
                record['currency'] = 'EUR'
            if len(errors) == before:
                key = identity(record)
                if key in seen:
                    errors.append(f'Row {row_number}: duplicate of row {seen[key]}')
                seen[key] = row_number
                records.append(record)
        if not records and not errors:
            errors.append('No tournaments found')
        if len(records) > 500:
            errors.append('Maximum 500 tournaments per import; split the workbook')
        if errors:
            raise ValidationError('\n'.join(errors))
        return records
    finally:
        wb.close()


def check_existing(records, db):
    keys = {identity(r): r['_row'] for r in records}
    errors = []
    for snapshot in db.collection('tournaments').stream():
        try:
            key = identity(snapshot.to_dict())
        except (KeyError, ValueError, TypeError, AttributeError) as exc:
            raise ValidationError(f'Cannot check existing tournament {snapshot.id}: invalid identity fields') from exc
        if key in keys:
            errors.append(f'Row {keys[key]}: duplicate of existing tournament {snapshot.id}')
    if errors:
        raise ValidationError('\n'.join(errors))


def connect(project, organizer):
    import firebase_admin
    from firebase_admin import auth, credentials, firestore
    app = firebase_admin.initialize_app(credentials.ApplicationDefault(), {'projectId': project})
    auth.get_user(organizer, app=app)
    return firestore.client(app=app)


def run_import(args, db=None):
    records = parse(args.file)
    print(f'Project: {args.project}\nOrganizer: {args.organizer_id}\nTournaments: {len(records)}')
    documents = [document(r, args.organizer_id) for r in records]
    for record, data in zip(records, documents):
        print(f'\nRow {record["_row"]}:\n' + json.dumps(data, ensure_ascii=False, indent=2))
    if db is None:
        db = connect(args.project, args.organizer_id)
    check_existing(records, db)
    if args.dry_run:
        print('Dry run successful; no writes performed.')
        return 0
    if not sys.stdin.isatty():
        print('Import requires an interactive terminal.', file=sys.stderr)
        return 1
    try:
        confirmed = input('Type IMPORT to create these tournaments: ') == 'IMPORT'
    except EOFError:
        confirmed = False
    if not confirmed:
        print('Import cancelled; no writes performed.')
        return 0
    check_existing(records, db)
    batch = db.batch()
    for data in documents:
        batch.create(db.collection('tournaments').document(data['id']), data)
    batch.commit()
    print(f'Created {len(documents)} tournaments:')
    for data in documents:
        print(data['id'])
    return 0


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command', required=True)
    create = sub.add_parser('template')
    create.add_argument('output', type=Path)
    importer = sub.add_parser('import')
    importer.add_argument('file', type=Path)
    importer.add_argument('-p', '--project', required=True)
    importer.add_argument('-oid', '--organizer-id', required=True)
    importer.add_argument('-n', '--dry-run', action='store_true')
    args = parser.parse_args()
    try:
        if args.command == 'template':
            template(args.output)
            print(f'Template written: {args.output}')
            return 0
        return run_import(args)
    except KeyboardInterrupt:
        print('Interrupted.', file=sys.stderr)
        return 130
    except Exception as exc:
        print(f'Error: {exc}', file=sys.stderr)
        return 1


if __name__ == '__main__':
    sys.exit(main())
